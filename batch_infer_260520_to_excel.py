import argparse
import base64
import json
from pathlib import Path
from typing import Any, Optional

import requests
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font
from PIL import Image as PILImage


BASE_URL = "http://127.0.0.1:30600"
MODEL_NAME = "PointerMeterInferModel"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
THUMBNAIL_MAX_WIDTH = 180
THUMBNAIL_MAX_HEIGHT = 120


def get_default_paths() -> tuple[Path, Path, Path]:
    api_dir = Path(__file__).resolve().parent
    image_dir = api_dir / "data"
    output_excel = api_dir / "data_infer_results.xlsx"
    result_image_dir = api_dir / "data_infer_result_images"
    return image_dir, output_excel, result_image_dir


def build_payload(image_path: Path, device: str) -> dict[str, Any]:
    return {
        "modelName": MODEL_NAME,
        "inferData": str(image_path.resolve()),
        "dataType": "filepath",
        "cameraTimeout": 10,
        "inferConfig": {
            "device": device,
            "result_pointer_image": True,
            "result_mask_image": True,
        },
    }


def to_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def decode_base64_image(image_base64: Optional[str], output_path: Path) -> Optional[Path]:
    if not image_base64:
        return None

    try:
        if "," in image_base64:
            image_base64 = image_base64.split(",", 1)[1]
        image_bytes = base64.b64decode(image_base64)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(image_bytes)
        return output_path
    except Exception as exc:
        print(f"[WARN] 图片base64解码失败: {output_path.name}, {exc}")
        return None


def fit_image_size(image_path: Path) -> tuple[int, int]:
    with PILImage.open(image_path) as image:
        width, height = image.size

    ratio = min(THUMBNAIL_MAX_WIDTH / width, THUMBNAIL_MAX_HEIGHT / height, 1)
    return int(width * ratio), int(height * ratio)


def add_image(ws, image_path: Optional[Path], cell: str) -> None:
    if image_path is None or not image_path.exists():
        return

    try:
        image = ExcelImage(str(image_path))
        image.width, image.height = fit_image_size(image_path)
        ws.add_image(image, cell)
    except Exception as exc:
        print(f"[WARN] Excel插入图片失败: {image_path}, {exc}")


def infer_one_image(base_url: str, image_path: Path, device: str, timeout: int) -> dict[str, Any]:
    payload = build_payload(image_path, device)
    response = requests.post(f"{base_url.rstrip('/')}/infer", json=payload, timeout=timeout)
    response.raise_for_status()
    return response.json()


def setup_sheet(ws) -> None:
    headers = [
        "文件名",
        "原图",
        "status",
        "message",
        "result",
        "result pointer image",
        "result mask image",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 34,
        "B": 26,
        "C": 12,
        "D": 52,
        "E": 18,
        "F": 26,
        "G": 26,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:G1"


def collect_images(image_dir: Path) -> list[Path]:
    return sorted(
        path for path in image_dir.iterdir()
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
    )


def write_result_row(
    ws,
    row_index: int,
    image_path: Path,
    response_json: dict[str, Any],
    pointer_image_path: Optional[Path],
    mask_image_path: Optional[Path],
) -> None:
    ws.cell(row=row_index, column=1, value=image_path.name)
    ws.cell(row=row_index, column=3, value=to_cell_text(response_json.get("status")))
    ws.cell(row=row_index, column=4, value=to_cell_text(response_json.get("message")))
    ws.cell(row=row_index, column=5, value=to_cell_text(response_json.get("result")))

    ws.row_dimensions[row_index].height = 96
    for column_index in range(1, 8):
        ws.cell(row=row_index, column=column_index).alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    add_image(ws, image_path, f"B{row_index}")
    add_image(ws, pointer_image_path, f"F{row_index}")
    add_image(ws, mask_image_path, f"G{row_index}")


def main() -> int:
    default_image_dir, default_output_excel, default_result_image_dir = get_default_paths()

    parser = argparse.ArgumentParser(description="批量调用指针表计识别API并导出Excel。")
    parser.add_argument("--base-url", default=BASE_URL, help="FastAPI服务地址。")
    parser.add_argument("--image-dir", type=Path, default=default_image_dir, help="待推理图片目录。")
    parser.add_argument("--output", type=Path, default=default_output_excel, help="输出Excel路径。")
    parser.add_argument(
        "--result-image-dir",
        type=Path,
        default=default_result_image_dir,
        help="接口返回图片的缓存目录。",
    )
    parser.add_argument("--device", default="0", help='inferConfig.device，默认使用"0"。')
    parser.add_argument("--timeout", type=int, default=120, help="单张图片请求超时时间，单位秒。")
    args = parser.parse_args()

    image_dir = args.image_dir.resolve()
    output_excel = args.output.resolve()
    result_image_dir = args.result_image_dir.resolve()

    if not image_dir.exists():
        raise FileNotFoundError(f"图片目录不存在: {image_dir}")

    images = collect_images(image_dir)
    if not images:
        raise FileNotFoundError(f"图片目录内未找到图片: {image_dir}")

    try:
        root_response = requests.get(f"{args.base_url.rstrip('/')}/", timeout=10)
        root_response.raise_for_status()
    except Exception as exc:
        print(f"[WARN] 服务根路径检查失败，将继续尝试调用/infer: {exc}")

    output_excel.parent.mkdir(parents=True, exist_ok=True)
    result_image_dir.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "infer_results"
    setup_sheet(ws)

    print(f"开始处理 {len(images)} 张图片，输入目录: {image_dir}")

    for index, image_path in enumerate(images, start=2):
        print(f"[{index - 1}/{len(images)}] 推理: {image_path.name}")
        pointer_image_path = None
        mask_image_path = None

        try:
            response_json = infer_one_image(args.base_url, image_path, args.device, args.timeout)
            pointer_image_path = decode_base64_image(
                response_json.get("result_pointer_image"),
                result_image_dir / f"{image_path.stem}_result_pointer.jpg",
            )
            mask_image_path = decode_base64_image(
                response_json.get("result_mask_image"),
                result_image_dir / f"{image_path.stem}_result_mask.jpg",
            )
        except Exception as exc:
            response_json = {
                "status": False,
                "message": f"请求或解析失败: {exc}",
                "result": None,
            }
            print(f"[ERROR] {image_path.name}: {exc}")

        write_result_row(ws, index, image_path, response_json, pointer_image_path, mask_image_path)

    wb.save(output_excel)
    print(f"Excel已保存: {output_excel}")
    print(f"返回图片缓存目录: {result_image_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
